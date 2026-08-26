"""Tests for workbook ordering and the applied-jobs sheet.

Seams under test (S7):
  - write_workbook(...) rebuilds each sheet's data region sorted latest-first
    (newest posted date at top, undated last), stable across runs.
  - Applied jobs are excluded from the deliverable sheet and listed in a
    dedicated 'Applied Jobs' sheet.
"""
from __future__ import annotations

from datetime import date, timedelta

from openpyxl import load_workbook

from scraper.data_model import ScrapedJob
from scraper.excel_writer import write_workbook
from scraper.utils import job_id


def _job(title: str, posted: date | None, scraped_at: str = "2026-08-24T09:00:00") -> ScrapedJob:
    url = f"https://example.com/{title.lower().replace(' ', '-')}"
    return ScrapedJob(
        job_id=job_id(url),
        source="grepjob",
        ats_url=url,
        ats_provider="greenhouse",
        title=title,
        company="Acme",
        location="Remote (Worldwide)",
        posted_date=posted,
        accepted=True,
        scraped_at=scraped_at,
    )


def _titles(path, sheet="Zero Experience Jobs", url_col=10, title_col=5):
    wb = load_workbook(path)
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(min_row=2):
        if row[url_col - 1].value:
            out.append(row[title_col - 1].value)
    return out


def test_newest_jobs_are_written_first_across_runs(tmp_path):
    path = tmp_path / "jobs.xlsx"
    d = date(2026, 8, 24)

    # First run writes an older job, then an older-still one
    write_workbook(
        path=path,
        qualifying_jobs=[_job("Old Job B", d - timedelta(days=5)),
                         _job("Old Job A", d - timedelta(days=9))],
        all_jobs=[], run_summary={"mode": "daily"},
    )
    # Second run contributes the freshest job
    write_workbook(
        path=path,
        qualifying_jobs=[_job("Fresh Job", d)],
        all_jobs=[], run_summary={"mode": "daily"},
    )

    assert _titles(path) == ["Fresh Job", "Old Job B", "Old Job A"]


def test_undated_jobs_sort_after_dated_ones(tmp_path):
    path = tmp_path / "jobs.xlsx"
    write_workbook(
        path=path,
        qualifying_jobs=[_job("No Date", None),
                         _job("Dated", date(2026, 8, 20))],
        all_jobs=[],
        run_summary={"mode": "daily"},
    )
    assert _titles(path) == ["Dated", "No Date"]


def test_applied_sheet_lists_applied_and_queue_excludes_them(tmp_path):
    path = tmp_path / "jobs.xlsx"
    fresh = _job("Want This", date(2026, 8, 23))
    applied = _job("Already Applied", date(2026, 8, 22))

    write_workbook(
        path=path,
        qualifying_jobs=[fresh, applied],
        all_jobs=[fresh, applied],
        run_summary={"mode": "daily"},
        excluded_job_ids={applied.job_id},
        applied_records=[{
            "applied_at": "2026-08-24T08:00:00",
            "source": applied.source,
            "title": applied.title,
            "company": applied.company,
            "location": applied.location,
            "ats_url": applied.ats_url,
            "note": "",
        }],
    )

    assert "Already Applied" not in _titles(path)
    assert "Already Applied" not in _titles(path, sheet="All Filtered Jobs")

    wb = load_workbook(path)
    assert "Applied Jobs" in wb.sheetnames
    ws = wb["Applied Jobs"]
    assert applied.ats_url in [c.value for c in ws["F"]]  # Apply URL column
    assert any(row[2].value == "Already Applied" for row in ws.iter_rows(min_row=2))


def test_audit_sheets_keep_everything_including_excluded(tmp_path):
    path = tmp_path / "jobs.xlsx"
    fresh = _job("Keep Visible", date(2026, 8, 23))
    applied = _job("Applied But Audited", date(2026, 8, 22))
    write_workbook(
        path=path,
        qualifying_jobs=[fresh, applied],
        all_jobs=[fresh, applied],
        run_summary={"mode": "daily"},
        all_scraped_jobs=[fresh, applied],
        excluded_job_ids={applied.job_id},
        applied_records=[{
            "applied_at": "2026-08-24T08:00:00",
            "title": applied.title,
            "ats_url": applied.ats_url,
        }],
    )
    assert "Applied But Audited" in _titles(path, sheet="All Scraped Jobs",
                                            url_col=9, title_col=4)
