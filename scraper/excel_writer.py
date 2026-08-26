"""
Excel writer using openpyxl.

Produces a single .xlsx workbook with FIVE sheets:
  - Sheet "Zero Experience Jobs":  the qualifying jobs (the deliverable queue)
  - Sheet "All Filtered Jobs":     every location+category-matched job (audit)
  - Sheet "All Scraped Jobs":      EVERY scraped job with filter outcome (audit)
  - Sheet "Applied Jobs":          jobs you marked applied via feedback
  - Sheet "Run Summary":           per-run stats and per-source breakdown

Every job sheet's data region is rebuilt on each run sorted latest-first
(newest posted date at top, undated last), so the freshest jobs are always
the first rows you see. Jobs whose id is in excluded_job_ids are dropped
from the deliverable and filtered sheets but kept in the full-audit sheet.
"""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .data_model import ScrapedJob

log = logging.getLogger("scraper.excel")


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ACCEPT_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")   # green
REJECT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")   # orange
ERROR_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")    # yellow
LOC_MATCH_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")  # light blue
FEEDBACK_FILL = PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid")  # purple

CELL_FONT = Font(name="Calibri", size=10)
CELL_ALIGN = Alignment(vertical="top", wrap_text=True)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


QUALIFYING_COLUMNS = [
    ("Date Scraped",        "scraped_at",         20),
    ("Source",              "source",             18),
    ("Location Type",       "location_type",      16),
    ("Posted Date",         "posted_date",        14),
    ("Job Title",           "title",              40),
    ("Company",             "company",            28),
    ("Location",            "location",           30),
    ("Seniority",           "seniority",          14),
    ("ATS Provider",        "ats_provider",       14),
    ("Apply URL",           "ats_url",            50),
    ("Category Matched",    "category_signals",   32),
    ("Experience Matched",  "experience_signals", 30),
    ("Years Mentioned",     "years_mentioned",    30),
    ("Filter Reason",       "filter_reason",      36),
    ("Description Preview", "description_preview", 80),
]

ALL_FILTERED_COLUMNS = [
    ("Date Scraped",        "scraped_at",         20),
    ("Source",              "source",             18),
    ("Location Type",       "location_type",      16),
    ("Posted Date",         "posted_date",        14),
    ("Job Title",           "title",              40),
    ("Company",             "company",            28),
    ("Location",            "location",           30),
    ("Seniority",           "seniority",          14),
    ("ATS Provider",        "ats_provider",       14),
    ("Apply URL",           "ats_url",            50),
    ("Loc Match",           "is_location_match",  10),
    ("Cat Match",           "is_category_match",  10),
    ("Accepted (0 Yrs)",    "accepted",           14),
    ("Filter Reason",       "filter_reason",      50),
    ("Description Preview", "description_preview", 60),
]

ALL_SCRAPED_COLUMNS = [
    ("Date Scraped",        "scraped_at",         20),
    ("Source",              "source",             18),
    ("Posted Date",         "posted_date",        14),
    ("Job Title",           "title",              40),
    ("Company",             "company",            28),
    ("Location",            "location",           30),
    ("Seniority",           "seniority",          14),
    ("ATS Provider",        "ats_provider",       14),
    ("Apply URL",           "ats_url",            50),
    ("Loc Match",           "is_location_match",  10),
    ("Loc Type",            "location_type",      16),
    ("Cat Match",           "is_category_match",  10),
    ("Accepted",            "accepted",           10),
    ("Filter Reason",       "filter_reason",      50),
    ("Error",               "error",              30),
    ("Description Preview", "description_preview", 50),
]

APPLIED_COLUMNS = [
    ("Date Applied",  "applied_at", 20),
    ("Source",        "source",     18),
    ("Job Title",     "title",      40),
    ("Company",       "company",    28),
    ("Location",      "location",   30),
    ("Apply URL",     "ats_url",    50),
    ("Note",          "note",       40),
]

SUMMARY_COLUMNS = [
    ("Metric", "metric", 32),
    ("Value",  "value",  50),
]

JOB_SHEETS = {
    "Zero Experience Jobs": QUALIFYING_COLUMNS,
    "All Filtered Jobs": ALL_FILTERED_COLUMNS,
    "All Scraped Jobs": ALL_SCRAPED_COLUMNS,
}
ALL_SHEETS = ("Zero Experience Jobs", "All Filtered Jobs", "All Scraped Jobs",
              "Applied Jobs", "Run Summary")


def write_workbook(
    path: Path,
    qualifying_jobs: List[ScrapedJob],
    all_jobs: List[ScrapedJob],
    run_summary: dict,
    all_scraped_jobs: Optional[List[ScrapedJob]] = None,
    excluded_job_ids: Optional[set] = None,
    applied_records: Optional[List[dict]] = None,
) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    excluded = excluded_job_ids or set()

    if path.exists():
        wb = load_workbook(path)
        _ensure_sheets(wb)
    else:
        wb = Workbook()
        _init_sheets(wb)

    # Merge stored rows with this run's jobs, dedupe by job id, then
    # rebuild each sheet latest-first.
    _rebuild_job_sheet(
        wb["Zero Experience Jobs"], new_jobs=qualifying_jobs,
        exclude_ids=excluded,
        highlight_fn=_hl_accepted,
    )
    _rebuild_job_sheet(
        wb["All Filtered Jobs"], new_jobs=all_jobs,
        exclude_ids=excluded,
        highlight_fn=_hl_filtered_row,
    )
    _rebuild_job_sheet(
        wb["All Scraped Jobs"], new_jobs=(all_scraped_jobs or []),
        exclude_ids=set(),
        highlight_fn=_hl_scraped_row,
    )
    _write_applied(wb["Applied Jobs"], applied_records or [])
    _write_summary(wb["Run Summary"], run_summary)

    tmp_path = path.with_suffix(".tmp.xlsx")
    wb.save(tmp_path)
    tmp_path.replace(path)
    log.info(
        "excel saved: %s (qualifying=%d filtered=%d scraped=%d applied=%d)",
        path, len(qualifying_jobs), len(all_jobs), len(all_scraped_jobs or []),
        len(applied_records or []),
    )


# --------------------------------------------------------------------------- #
#  Row model: every sheet body is a list of flat dicts keyed by column attr
# --------------------------------------------------------------------------- #

def _job_to_row(job: ScrapedJob, columns) -> dict:
    return {attr: _cell_value(job, attr) for _h, attr, _w in columns}


def _read_rows(ws, columns) -> List[dict]:
    """Existing sheet body as row dicts (values already formatted)."""
    rows: List[dict] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in values):
            continue
        row = {}
        for (_h, attr, _w), v in zip(columns, values):
            row[attr] = v
        rows.append(row)
    return rows


def _row_id(row: dict) -> str:
    from .utils import job_id as compute_id
    url = str(row.get("ats_url") or "")
    return compute_id(url) if url else ""


def _sort_key(row: dict):
    from datetime import datetime as _dt
    posted = row.get("posted_date")
    if isinstance(posted, _dt):
        posted = posted.date()
    if isinstance(posted, str):
        try:
            posted = date.fromisoformat(str(posted)[:10])
        except ValueError:
            posted = None
    return (posted is not None, posted or date.min, str(row.get("scraped_at") or ""))


def _rebuild_job_sheet(ws, new_jobs: Iterable[ScrapedJob], exclude_ids: set,
                       highlight_fn) -> None:
    columns = JOB_SHEETS[ws.title]
    rows = _read_rows(ws, columns)
    rows.extend(_job_to_row(j, columns) for j in new_jobs)

    deduped: List[dict] = []
    seen: set = set()
    for row in rows:
        rid = _row_id(row)
        if rid and (rid in seen or rid in exclude_ids):
            continue
        if rid:
            seen.add(rid)
        deduped.append(row)
    deduped.sort(key=_sort_key, reverse=True)

    _clear_body(ws)
    _style_header_row(ws, columns)
    _write_rows(ws, deduped, columns, highlight_fn)


def _clear_body(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def _write_rows(ws, rows: List[dict], columns, highlight_fn) -> None:
    for i, row in enumerate(rows):
        row_idx = i + 2
        fill = highlight_fn(row)
        for col_idx, (_header, attr, _width) in enumerate(columns, start=1):
            value = row.get(attr)
            value = "" if value is None else value
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGN
            cell.fill = fill
            cell.border = THIN_BORDER
            if attr == "ats_url" and value:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, target=str(value))
                cell.font = LINK_FONT
        preview_len = len(str(row.get("description_preview") or ""))
        ws.row_dimensions[row_idx].height = min(120, max(20, preview_len // 6))


def _truthy(value) -> bool:
    return value in (True, "Yes", "yes", "Y")


def _hl_accepted(row: dict):
    return ACCEPT_FILL


def _hl_filtered_row(row: dict):
    return ACCEPT_FILL if _truthy(row.get("accepted")) else REJECT_FILL


def _hl_scraped_row(row: dict):
    if _truthy(row.get("accepted")):
        return ACCEPT_FILL
    if str(row.get("filter_reason") or "").startswith("feedback:"):
        return FEEDBACK_FILL
    if row.get("error"):
        return ERROR_FILL
    if _truthy(row.get("is_location_match")):
        return LOC_MATCH_FILL
    return REJECT_FILL


def _write_applied(ws, records: List[dict]) -> None:
    rows = _read_rows(ws, APPLIED_COLUMNS)
    known_urls = {str(r.get("ats_url") or "") for r in rows}
    for rec in records:
        row = {attr: rec.get(attr, "") for _h, attr, _w in APPLIED_COLUMNS}
        if str(row.get("ats_url") or "") in known_urls:
            continue
        known_urls.add(str(row.get("ats_url") or ""))
        rows.append(row)
    rows.sort(key=lambda r: str(r.get("applied_at") or ""), reverse=True)
    _clear_body(ws)
    _style_header_row(ws, APPLIED_COLUMNS)
    _write_rows(ws, rows, APPLIED_COLUMNS, lambda row: ACCEPT_FILL)


# --------------------------------------------------------------------------- #
#  Sheets / styling helpers
# --------------------------------------------------------------------------- #

def _ensure_sheets(wb: Workbook) -> None:
    for name in ALL_SHEETS:
        if name not in wb.sheetnames:
            wb.create_sheet(name)
            _style_header_row(wb[name], _columns_for(name))


def _init_sheets(wb: Workbook) -> None:
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for name in ALL_SHEETS:
        wb.create_sheet(name)
        _style_header_row(wb[name], _columns_for(name))


def _columns_for(sheet_name: str):
    if sheet_name in JOB_SHEETS:
        return JOB_SHEETS[sheet_name]
    if sheet_name == "Applied Jobs":
        return APPLIED_COLUMNS
    return SUMMARY_COLUMNS


def _style_header_row(ws, columns) -> None:
    for col_idx, (header, _attr, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "C2"


def _cell_value(job: ScrapedJob, attr: str):
    if attr == "scraped_at":
        return job.scraped_at or ""
    if attr == "location_type":
        return job.location_type or ""
    if attr == "seniority":
        return job.seniority or ""
    if attr == "posted_date":
        return job.posted_date or ""
    if attr == "is_location_match":
        return "Yes" if job.is_location_match else "No"
    if attr == "is_category_match":
        return "Yes" if job.is_category_match else "No"
    if attr == "accepted":
        return "Yes" if job.accepted else "No"
    if attr == "category_signals":
        return "; ".join(job.category_signals) if job.category_signals else ""
    if attr == "experience_signals":
        return "; ".join(job.experience_signals) if job.experience_signals else ""
    if attr == "years_mentioned":
        return "; ".join(job.years_mentioned) if job.years_mentioned else ""
    if attr == "description_preview":
        return job.description_preview
    if attr == "error":
        return job.error or ""
    return getattr(job, attr, "") or ""


def _write_summary(ws, summary: dict) -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.font = CELL_FONT
            cell.border = Border()
    _style_header_row(ws, SUMMARY_COLUMNS)
    rows = [
        ("Run timestamp",            summary.get("run_at", "")),
        ("Run mode",                 summary.get("mode", "")),
        ("Days back",                summary.get("days_back", "")),
        ("Sources scraped",          ", ".join(summary.get("sources_scraped", []))),
        ("Jobs parsed (total)",      summary.get("parsed_count", 0)),
    ]
    by_source = summary.get("by_source", {})
    for src, count in sorted(by_source.items()):
        rows.append((f"  - {src}", count))
    rows.extend([
        ("Location matches",         summary.get("location_match_count", 0)),
        ("Category matches",         summary.get("category_match_count", 0)),
        ("Zero-experience jobs",     summary.get("qualifying_count", 0)),
        ("Rejected (location)",      summary.get("not_location_count", 0)),
        ("Rejected (category)",      summary.get("wrong_category_count", 0)),
        ("Rejected (experience)",    summary.get("has_experience_count", 0)),
        ("Rejected (your feedback)", summary.get("feedback_rejected_count", 0)),
        ("Skipped (already applied)", summary.get("applied_seen_count", 0)),
        ("Errors",                   summary.get("error_count", 0)),
        ("Newly written",            summary.get("new_written", 0)),
        ("Duration (seconds)",       round(summary.get("duration_seconds", 0.0), 2)),
    ])
    for i, (metric, value) in enumerate(rows, start=2):
        c1 = ws.cell(row=i, column=1, value=metric)
        c1.font = Font(name="Calibri", bold=True, size=10)
        c1.border = THIN_BORDER
        c2 = ws.cell(row=i, column=2, value=value)
        c2.font = CELL_FONT
        c2.border = THIN_BORDER
